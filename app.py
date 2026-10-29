import re
from collections import Counter
from io import BytesIO
from pathlib import Path
import streamlit as st
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

st.set_page_config(page_title="StudyMate Local", page_icon="📚", layout="wide")
st.markdown("""
<style>
.hero{padding:1.3rem 1.5rem;border-radius:18px;background:#172554;color:white;margin-bottom:1rem}
.hero h1{margin:0;color:white}.hero p{color:#dbeafe;margin:.4rem 0 0}
div.stButton>button{border-radius:10px;font-weight:600}
</style>
<div class="hero"><h1>📚 StudyMate Local</h1><p>Document study assistant — no API key, no cloud AI service.</p></div>
""", unsafe_allow_html=True)

STOP = set("a an the and or but if then than so because as at by for from in into of on to up with about is are was were be been being have has had do does did can could should would may might will this that these those it its they them their he she we you your i me my our not also some any each all both few many much which who what when where why how".split())
EXTS = ["pdf", "docx", "pptx", "txt", "md"]

def extract(file):
    ext = Path(file.name).suffix.lower()
    data = file.getvalue()
    if ext == ".pdf":
        return "\n\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages)
    if ext == ".docx":
        d = Document(BytesIO(data))
        out = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            out.extend(" | ".join(c.text.strip() for c in r.cells) for r in t.rows)
        return "\n".join(out)
    if ext == ".pptx":
        prs = Presentation(BytesIO(data))
        return "\n\n".join(f"Slide {i}: " + "\n".join(s.text for s in slide.shapes if hasattr(s,"text") and s.text.strip()) for i,slide in enumerate(prs.slides,1))
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    return ""

def sentences(text):
    return [x.strip(" \n\t•-") for x in re.split(r"(?<=[.!?])\s+|\n+", text) if len(x.strip()) > 25]

def tokens(text):
    return re.findall(r"\b[\w'-]+\b", text.lower())

def terms(text, n=25):
    return [w for w,c in Counter(w for w in tokens(text) if w not in STOP and len(w)>3 and not w.isdigit()).most_common(n)]

def summary(text, n=9):
    ss = sentences(text)
    if not ss: return "No readable text found. Scanned PDFs need OCR before upload."
    freq = Counter(w for w in tokens(text) if w not in STOP and len(w)>2)
    ranked = []
    for i,s in enumerate(ss):
        ws = [w for w in tokens(s) if w not in STOP and len(w)>2]
        score = sum(freq[w] for w in ws)/max(1,len(ws))
        if re.search(r"\b(definition|defined|means|refers to|important|advantages|disadvantages|types|process|example|therefore)\b",s,re.I): score *= 1.15
        ranked.append((score,i,s))
    return "\n\n".join("• "+s for _,_,s in sorted(sorted(ranked,reverse=True)[:n],key=lambda x:x[1]))

def short_questions(text, n=15):
    ss=sentences(text); out=[]
    for s in ss:
        m=re.match(r"^(.{3,70}?)\s+(?:is|are|means|refers to)\s+(.+)$",s,re.I)
        if m and 1 < len(m.group(1).split()) < 9:
            out.append((f"What is {m.group(1).strip(' ,:;')}?",s))
        if len(out)>=n: break
    for term in terms(text, n*3):
        if len(out)>=n: break
        evidence=next((s for s in ss if re.search(r"\b"+re.escape(term)+r"\b",s,re.I)),None)
        if evidence and not any(term.lower() in q.lower() for q,a in out):
            out.append((f"Explain the term '{term}'.", evidence))
    return out

def long_questions(text, n=10):
    ss=sentences(text); out=[]
    templates=["Explain {t} in detail, including its key features and significance.","Discuss {t} with suitable examples from the material.","Describe the main concepts related to {t} and how they work.","Write a detailed note on {t}, including its applications and limitations."]
    for i,t in enumerate(terms(text,n*2)):
        ev=[s for s in ss if re.search(r"\b"+re.escape(t)+r"\b",s,re.I)]
        out.append((templates[i%len(templates)].format(t=t), " ".join(ev[:3]) or "Insufficient supporting text."))
        if len(out)>=n: break
    return out

def mcqs(text,n=12):
    ss=sentences(text); ts=terms(text,n*4); out=[]
    for t in ts:
        s=next((s for s in ss if re.search(r"\b"+re.escape(t)+r"\b",s,re.I) and len(s)<260),None)
        distractors=[x for x in ts if x.lower()!=t.lower() and x.lower() not in (s or "").lower()][:3]
        if not s or len(distractors)<3: continue
        q=re.sub(r"\b"+re.escape(t)+r"\b","________",s,count=1,flags=re.I)
        opts=[t]+distractors; shift=len(out)%4; opts=opts[shift:]+opts[:shift]
        out.append((q,opts,"ABCD"[opts.index(t)],t))
        if len(out)>=n: break
    return out

def answer(text, question):
    ts=[w for w in tokens(question) if w not in STOP and len(w)>2]
    ranked=[]
    for i,s in enumerate(sentences(text)):
        overlap=sum(w in set(tokens(s)) for w in ts)
        if overlap: ranked.append((overlap/max(1,len(ts))**.5,i,s))
    return "\n\n".join("• "+s for _,_,s in sorted(sorted(ranked,reverse=True)[:5],key=lambda x:x[1])) or "No close match found. Try a more specific question or keywords."

def generate(kind,text,question):
    if kind=="Summary":
        return "## Summary\n\n"+summary(text)+"\n\n## Key terms\n\n"+", ".join(terms(text,20))
    if kind=="Short questions":
        q=short_questions(text)
        return "## Short questions and source-based answers\n\n"+"\n\n".join(f"**{i}. {a}**\n\nAnswer: {b}" for i,(a,b) in enumerate(q,1)) or "Not enough structured text to create questions."
    if kind=="Long questions":
        q=long_questions(text)
        return "## Long questions with supporting passages\n\n"+"\n\n".join(f"**{i}. {a}**\n\nSuggested points: {b}" for i,(a,b) in enumerate(q,1))
    if kind=="MCQ quiz":
        q=mcqs(text)
        return "## MCQ quiz\n\n"+"\n\n".join(f"**{i}. {a}**\n\n"+"\n".join(f"{'ABCD'[j]}. {o}" for j,o in enumerate(opts))+f"\n\n**Answer: {c}. {t}**" for i,(a,opts,c,t) in enumerate(q,1)) or "Not enough suitable text to create MCQs."
    if kind=="Flashcards":
        cards=[]
        for t in terms(text,20):
            evidence=next((s for s in sentences(text) if re.search(r"\\b"+re.escape(t)+r"\\b",s,re.I)),"No definition found.")
            cards.append(f"**Q: What should you remember about {t}?**\\n\\nA: {evidence}")
        return "## Flashcards\\n\\n"+"\\n\\n".join(cards)
    if kind=="Important topics":
        return "## Topics to review\n\n"+"\n".join(f"{i}. **{t.title()}** — appears frequently in the document." for i,t in enumerate(terms(text,20),1))+"\n\nThese are frequency-based suggestions, not exam predictions."
    return "## Answer from matching passages\n\n"+answer(text,question)

if "docs" not in st.session_state: st.session_state.docs={}
if "results" not in st.session_state: st.session_state.results=[]
with st.sidebar:
    st.header("📂 Upload files")
    uploads=st.file_uploader("PDF, Word, PowerPoint, TXT, Markdown",type=EXTS,accept_multiple_files=True)
    if uploads:
        for f in uploads:
            if f.name not in st.session_state.docs:
                try:
                    txt=extract(f).replace("\x00"," ").strip()
                    if txt: st.session_state.docs[f.name]=txt
                    else: st.warning(f"No selectable text found in {f.name}. It may be scanned.")
                except Exception as e: st.error(f"Could not read {f.name}: {e}")
    names=list(st.session_state.docs)
    selected=st.multiselect("Documents to use",names,default=names)
    if st.button("Clear documents"):
        st.session_state.docs={}; st.session_state.results=[]; st.rerun()
    st.caption("All processing is local. No API key is needed.")
text="\n\n".join("FILE: "+n+"\n"+st.session_state.docs[n] for n in selected)
st.subheader("Turn documents into study material")
if not text:
    st.info("Upload a document in the sidebar to begin.")
else:
    col1,col2=st.columns([1,1])
    with col1: kind=st.selectbox("Study tool",["Summary","Short questions","Long questions","MCQ quiz","Flashcards","Important topics","Custom question"])
    with col2: question=st.text_input("Your question",placeholder="Example: Explain normalization")
    if st.button("Generate",type="primary",use_container_width=True):
        st.session_state.results.insert(0,(kind,generate(kind,text,question)))
    if st.session_state.results:
        st.divider(); st.subheader("Generated material")
        if st.button("Clear results"): st.session_state.results=[]; st.rerun()
        for i,(kind,result) in enumerate(st.session_state.results):
            with st.expander(kind,expanded=(i==0)):
                st.markdown(result)
                st.download_button("Download as .txt",result,file_name=f"studymate_{i+1}.txt",mime="text/plain",key=f"download_{i}")
st.divider()
st.caption("This no-key edition uses local rules and keyword matching, not a generative language model. Results may be less fluent or miss context; verify answers against your source.")
